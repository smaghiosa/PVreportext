"""
db_export.py

Exports the CURRENT state of the database back into the same multi-
sheet Excel format bulk_extract.py produces from fresh PDF extraction -
so a report can be regenerated for any tag/project already in the
database without re-processing the original PDFs. This is the
"database is the master repository, Excel is just a transfer/report
format" direction from the original project brief.

Usage:
    python db_export.py pvelite_data.db Master_Element_Data_export.xlsx
    python db_export.py pvelite_data.db out.xlsx --project ADVP
"""
import sys
import sqlite3
import argparse
import openpyxl
from openpyxl.styles import Font

from bulk_extract import _style_sheet


def _connect(db_path):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def _matching_equipment_ids(conn, project=None, category=None, tag_no=None):
    sql = "SELECT id FROM equipment_master WHERE 1=1"
    params = []
    if project:
        sql += " AND project = ?"
        params.append(project)
    if category:
        sql += " AND category = ?"
        params.append(category)
    if tag_no:
        sql += " AND tag_no = ?"
        params.append(tag_no)
    return [row["id"] for row in conn.execute(sql, params)]


def export_to_excel(db_path, output_path, project=None, category=None, tag_no=None, log=print):
    conn = _connect(db_path)
    equipment_ids = _matching_equipment_ids(conn, project, category, tag_no)
    if not equipment_ids:
        log("No matching records found in the database - nothing to export.")
        conn.close()
        return False
    log(f"Exporting {len(equipment_ids)} record(s) from the database...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # --- Master_Element_Data ---
    ws = wb.create_sheet("Master_Element_Data")
    cols = ["Project", "Category", "Equipment", "Element No", "Element Name", "Component",
            "Material", "Class", "UNS Number", "Diameter (m)", "Element Length (mm)",
            "Minimum Thickness (mm)", "Nominal Thickness (mm)", "Corrosion Allowance (mm)",
            "External Pressure (MPa)", "Design Pressure (MPa)", "Design Temperature (degC)",
            "Allowable Tensile Stress (MPa)", "Allowable Compressive Stress (MPa)",
            "Surface Area (m2)", "Inside Surface Area (m2)", "Weight (MT)", "Total Ele. Empty Wgt. (MT)"]
    ws.append(cols)
    q = """
        SELECT e.project, e.category, e.tag_no, d.element_no, d.element_name, d.component,
               d.material, d.class, d.uns_number, d.diameter_m, d.element_length_mm,
               d.minimum_thickness_mm, d.nominal_thickness_mm, d.corrosion_allowance_mm,
               d.external_pressure_mpa, d.design_pressure_mpa, d.design_temperature_c,
               d.allowable_tensile_stress_mpa, d.allowable_compressive_stress_mpa,
               d.surface_area_m2, d.inside_surface_area_m2, d.weight_mt, d.total_ele_empty_wgt_mt
        FROM element_data d JOIN equipment_master e ON e.id = d.equipment_id
        WHERE d.equipment_id IN ({}) ORDER BY e.project, e.tag_no, d.element_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        ws.append(list(row))

    # --- Loads_Foundation_Support ---
    ws = wb.create_sheet("Loads_Foundation_Support")
    cols = ["Project", "Category", "Tag No", "Name", "ASME Code", "Division", "Code Edition Year",
            "MDMT (degC)", "Wind Design Code", "Seismic Design Code", "Basic Wind Speed (Km/hr)",
            "Site Class", "Seismic Ss", "Seismic S1", "Seismic SDS", "Seismic SD1",
            "Seismic Response Coeff. Cs", "Max Wind Pressure (Kgs/m2)", "Wind Shear (Kgf)",
            "Earthquake Shear (Kgf)", "Wind Moment (Kgf-m)", "Earthquake Moment (Kgf-m)",
            "Fabricated Weight (MT)", "Erected Weight (MT)", "Operating Weight (MT)",
            "Shop Test Weight (MT)", "Bolt MOC", "Bolt Nominal Dia (mm)", "Bolt Circle Dia (mm)",
            "Number of Bolts"]
    ws.append(cols)
    q = """
        SELECT project, category, tag_no, equipment_name, asme_code, division, code_edition_year,
               mdmt_c, wind_code, seismic_code, basic_wind_speed_kmh, site_class, seismic_ss,
               seismic_s1, seismic_sds, seismic_sd1, seismic_response_coeff_cs,
               max_wind_pressure_kgm2, wind_shear_kgf, eq_shear_kgf, wind_moment_kgfm,
               eq_moment_kgfm, fabricated_mt, erected_mt, operating_mt, shop_test_mt, bolt_moc,
               bolt_nominal_dia_mm, bolt_circle_dia_mm, bolt_qty
        FROM equipment_master WHERE id IN ({}) ORDER BY project, tag_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        ws.append(list(row))

    # --- Work_Volume ---
    ws = wb.create_sheet("Work_Volume")
    cols = ["Project", "Category", "Tag No", "Name", "TL-TL (mm)", "Max Diameter (mm)",
            "Min Diameter (mm)", "L/D Ratio", "Erection Weight (MT)", "Platform Qualified",
            "Total Platform Area (m2)", "Insulation Qualified", "Insulation Area (m2)",
            "Painting Area (m2)", "FP Qualified", "FP Area (m2)"]
    ws.append(cols)
    q = """
        SELECT project, category, tag_no, equipment_name, tl_tl_mm, max_diameter_mm,
               min_diameter_mm, l_over_d_ratio, erected_mt,
               CASE has_platform WHEN 1 THEN 'Yes' ELSE 'No' END,
               total_platform_area_m2,
               CASE has_insulation WHEN 1 THEN 'Yes' ELSE 'No' END,
               insulation_area_m2, painting_area_m2,
               CASE has_fp WHEN 1 THEN 'Yes' ELSE 'No' END,
               fp_area_m2
        FROM equipment_master WHERE id IN ({}) ORDER BY project, tag_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        ws.append(list(row))

    # --- Platform_List ---
    ws = wb.create_sheet("Platform_List")
    cols = ["Project", "Category", "Tag No", "Element Name", "Diameter (mm)", "Detail Type",
            "Detail ID", "Start Angle (deg)", "End Angle (deg)", "Platform Width (mm)",
            "Platform Length (mm)", "Platform Area (m2)"]
    ws.append(cols)
    q = """
        SELECT e.project, e.category, e.tag_no, p.element_name, p.diameter_mm, 'Platform',
               p.detail_id, p.start_angle_deg, p.end_angle_deg, p.width_mm, p.length_mm, p.area_m2
        FROM platform_data p JOIN equipment_master e ON e.id = p.equipment_id
        WHERE p.equipment_id IN ({}) ORDER BY e.project, e.tag_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        ws.append(list(row))

    # --- weight_summ ---
    ws = wb.create_sheet("weight_summ")
    cols = ["Project", "Category", "Tag No", "Components", "Fabricated", "Shop Test",
            "Shipping", "Erected", "Empty", "Operating"]
    ws.append(cols)
    q = """
        SELECT e.project, e.category, e.tag_no, w.component, w.fabricated_kg, w.shop_test_kg,
               w.shipping_kg, w.erected_kg, w.empty_kg, w.operating_kg
        FROM weight_summation w JOIN equipment_master e ON e.id = w.equipment_id
        WHERE w.equipment_id IN ({}) ORDER BY e.project, e.tag_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        # "..." was normalized to NULL on the way into the database (the
        # correct DB representation) - restored here purely for visual
        # consistency with the original PDF-sourced export.
        ws.append([v if v is not None else ("..." if i >= 4 else None) for i, v in enumerate(row)])

    # --- Equipment_Summary ---
    ws = wb.create_sheet("Equipment_Summary")
    cols = ["Project", "Category", "Tag No", "Name", "File Kind", "Warnings", "Imported By", "Imported At"]
    ws.append(cols)
    q = """
        SELECT project, category, tag_no, equipment_name, file_kind, warnings_count,
               imported_by, imported_at
        FROM equipment_master WHERE id IN ({}) ORDER BY project, tag_no
    """.format(",".join("?" * len(equipment_ids)))
    for row in conn.execute(q, equipment_ids):
        ws.append(list(row))

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for c in ws[1]:
            c.font = Font(bold=True)
        _style_sheet(ws)

    wb.save(output_path)
    conn.close()
    log(f"Exported to {output_path}")
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("db_path")
    parser.add_argument("output_path")
    parser.add_argument("--project", default=None)
    parser.add_argument("--category", default=None)
    parser.add_argument("--tag", default=None)
    args = parser.parse_args()
    export_to_excel(args.db_path, args.output_path, args.project, args.category, args.tag)
